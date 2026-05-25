"""Helpers for extracting text from uploaded office documents."""
from __future__ import annotations

import csv
import io
import os

ALLOWED_UPLOAD_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".csv"}
MAX_UPLOAD_SIZE = 10 * 1024 * 1024  # 10MB


def extract_text_from_file(filename: str, content: bytes) -> str:
    """Extract text from uploaded document based on file type."""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        from PyPDF2 import PdfReader

        reader = PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages).strip()

    if ext == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(content))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

    if ext == ".pptx":
        from pptx import Presentation

        prs = Presentation(io.BytesIO(content))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
        return "\n\n".join(t for t in texts if t.strip())

    if ext == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheets_text = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells):
                    rows.append("\t".join(cells))
            if rows:
                sheets_text.append(f"[Sheet: {sheet_name}]\n" + "\n".join(rows))
        wb.close()
        return "\n\n".join(sheets_text)

    if ext == ".csv":
        reader = csv.reader(io.StringIO(content.decode("utf-8", errors="replace")))
        rows = ["\t".join(row) for row in reader]
        return "\n".join(rows)

    raise ValueError(f"Unsupported file type: {ext}")
